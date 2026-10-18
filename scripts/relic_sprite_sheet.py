"""Marks in ``docs/reliquias.xlsx`` which relics have their own sprite (dev tool, needs openpyxl).

Adds the column "Sprite propio" to the "Reliquias" sheet if missing and writes "Sí" for every
relic listed in ``assets/relics/relics.json`` and "No" for the rest:

    python scripts/relic_sprite_sheet.py
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).parent.parent
SHEET = ROOT / "docs" / "reliquias.xlsx"
INDEX = ROOT / "assets" / "relics" / "relics.json"
HEADER = "Sprite propio"


def main() -> int:
    done = json.loads(INDEX.read_text(encoding="utf-8")) if INDEX.is_file() else {}
    wb = openpyxl.load_workbook(SHEET)
    ws = wb["Reliquias"]
    head_row = 4
    col = None
    for c in range(1, ws.max_column + 1):
        if ws.cell(head_row, c).value == HEADER:
            col = c
    if col is None:
        col = ws.max_column + 1
        src = ws.cell(head_row, 1)
        cell = ws.cell(head_row, col, HEADER)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=src.fill.fgColor.rgb if src.fill and src.fill.fgColor else "2F3B5C")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 13
    thin = Side(style="thin", color="BFBFBF")
    yes = 0
    for r in range(head_row + 1, ws.max_row + 1):
        name = ws.cell(r, 2).value
        if not name:
            continue
        own = name in done
        yes += own
        cell = ws.cell(r, col, "Sí" if own else "No")
        cell.font = Font(name="Arial", size=10, bold=own, color="1E7B34" if own else "808080")
        cell.alignment = Alignment(horizontal="center", vertical="top")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb.save(SHEET)
    print(f"Sprite propio: {yes} relics")
    return yes


if __name__ == "__main__":
    main()
